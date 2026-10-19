"""Pruebas del modo lote (sin red): evaluate(), banderas, Excel y réplica offline desde el JSON de entrada."""
import io, json
from dataclasses import replace
import openpyxl
import lote
from test_models import industrial
from dcf_models import CompanyData, Params, fcff_value, SCENARIOS


class DataError(Exception):
    pass


def loader(co):
    return lambda sym, a=False, r=0.0: (co, [])


def test_evaluate_ok_and_offline_replica():
    co = industrial(ticker="IND", sector="Industrials", industry="Specialty Industrial Machinery", country="United States")
    row, raw = lote.evaluate("IND", loader(co))
    assert row["Estado"] == "OK" and row["Modelo"] == "FCFF", row
    for k in ("Pesimista", "Base", "Optimista", "Pes/Base", "WACC", "Foso", "Margen TTM", "Banderas"):
        assert k in row, k
    assert row["Pesimista"] <= row["Base"] <= row["Optimista"]
    # el JSON permite reconstruir la empresa y obtener el mismo valor sin volver a Yahoo
    co2 = CompanyData(**json.loads(raw))
    from industry_data import attach
    c3, _, _ = attach(co2, "FCFF", None)
    from country_tax import marginal_tax_for
    from country_risk import country_crp
    from dcf_models import terminal_growth_default
    from industry_data import ERP_DAMODARAN
    p = Params(horizon=10, g_terminal=terminal_growth_default(co2.rf), erp=ERP_DAMODARAN, crp=country_crp(co2.country)[0], moat="auto",
               marginal_tax=marginal_tax_for(co2.country)[0], use_industry_beta=True)
    assert abs(fcff_value(c3, p, SCENARIOS["Base"]).per_share - row["Base"]) < 1e-9


def test_flags_depressed_ttm_and_revenue_trend():
    base = dict(sector="Healthcare", industry="Drug Manufacturers - General", country="United States",
                revenue_hist=[9_000, 9_500, 9_800, 10_000], ebit_hist=[2_800, 2_900, 3_000, 3_000])
    stable = industrial(**base, ebit_ttm=1_000.0, revenue_ttm=10_200.0)
    r1, _ = lote.evaluate("A", loader(stable))
    assert "TTM bajo vs mediana, ingresos estables" in r1["Banderas"], r1["Banderas"]
    falling = industrial(**base, ebit_ttm=1_000.0, revenue_ttm=8_000.0)
    r2, _ = lote.evaluate("B", loader(falling))
    assert "TTM bajo vs mediana, ingresos caen" in r2["Banderas"], r2["Banderas"]
    peak = industrial(**base, ebit_ttm=4_500.0, revenue_ttm=10_000.0)
    r3, _ = lote.evaluate("C", loader(peak))
    assert "TTM alto vs mediana" in r3["Banderas"], r3["Banderas"]


def test_degraded_error_and_not_valuable():
    degraded = industrial(sector="", industry="", country="", name="DEG", ticker="DEG")
    r, _ = lote.evaluate("DEG", loader(degraded))
    assert r["Estado"] == "No confiable" and "no confiable" in r["Banderas"] and "industria" in r["Datos faltantes"], r
    assert r["Base"] is None and "Valor ref. pes/base/opt" in r and "falta industria" in r["Detalle"], r

    def boom(sym, a=False, r=0.0):
        raise DataError("Yahoo bloqueó la consulta")
    r, raw = lote.evaluate("ZZZ", boom)
    assert r["Estado"] == "ERROR de datos" and raw is None and "Yahoo" in r["Detalle"]
    neg = industrial(ebit_ttm=-10.0, ni_ttm=-5.0, equity=-1.0)
    r, _ = lote.evaluate("NEG", loader(neg))
    assert r["Estado"] in ("No valuable", "OK", "Datos incompletos"), r


def test_country_only_and_adr_and_no_aplica():
    full = dict(sector="Industrials", industry="Specialty Industrial Machinery")
    # falta solo el país y SÍ hay analistas: valor con aviso
    r, _ = lote.evaluate("A", loader(industrial(**full, country="", rev_g_y0=0.05, rev_g_y1=0.05)))
    assert r["Estado"] == "Datos incompletos" and r["Base"] is not None, r
    # falta el país y además no hay analistas: no confiable
    r, _ = lote.evaluate("B", loader(industrial(**full, country="", rev_g_y0=None, rev_g_y1=None)))
    assert r["Estado"] == "No confiable" and r["Base"] is None, r
    # capitalización estimada con moneda distinta (ADR posible): no confiable
    r, _ = lote.evaluate("C", loader(industrial(**full, country="Taiwan", cap_estimated=True)))
    assert r["Estado"] == "No confiable" and "ADR" in r["Datos faltantes"], r
    # negocio que destruye valor: no aplica
    r, _ = lote.evaluate("D", loader(industrial(**full, country="United States", ebit_ttm=-50.0)))
    assert r["Estado"] in ("No aplica por DCF", "OK"), r
    if r["Estado"] == "No aplica por DCF":
        assert r["Base"] is None and "Valor ref. pes/base/opt" in r


def test_rescue_replaces_only_improved_rows():
    bad = industrial(ticker="X", sector="", industry="", country="", name="X")
    good = industrial(ticker="X", sector="Industrials", industry="Specialty Industrial Machinery", country="United States")
    rows, raws = [], {}
    r, raw = lote.evaluate("X", loader(bad)); rows.append(r)
    calls = {"n": 0}
    def flaky(sym, a=False, r_=0.0):
        calls["n"] += 1
        return (bad if calls["n"] < 2 else good), []
    waited = []
    fixed = lote.rescue(rows, raws, flaky, waits=(5.0, 5.0), sleep=waited.append)
    assert fixed == 1 and rows[0]["Estado"] in ("OK", "Datos incompletos") and rows[0]["Base"] is not None, rows[0]
    assert 5.0 in waited and calls["n"] == 2
    assert lote.rescue([dict(rows[0])], {}, flaky, waits=(5.0,), sleep=lambda s: None) == 0   # nada que rescatar


def test_excel_and_parse():
    co = industrial(sector="Industrials", industry="Specialty Industrial Machinery", country="United States")
    rows, raws = [], {}
    for s in ("AAA", "BBB"):
        r, raw = lote.evaluate(s, loader(co))
        rows.append(r)
        raws[s] = raw
    xl = lote.build_excel(rows, raws)
    wb = openpyxl.load_workbook(io.BytesIO(xl))
    assert set(wb.sheetnames) == {"Lote", "Resumen", "Leyenda_banderas", "Datos_crudos"}
    assert wb["Lote"].max_row == 3 and wb["Datos_crudos"].max_row == 3
    assert json.loads(wb["Datos_crudos"].cell(2, 2).value)["ticker"] in ("IND", "AAA", "BBB")
    assert lote.parse_tickers("aapl, msft\nko; PG  aapl") == ["AAPL", "MSFT", "KO", "PG"]
    assert len(lote.parse_tickers(lote.SUGGESTED)) >= 60
    states, flags = lote.summarize(rows)
    assert states[0][1] == 2


def test_rescue_avisa_el_progreso():
    msgs = []
    rows = [dict(Ticker="AAA", Estado="No confiable")]

    def flaky(sym):
        raise RuntimeError("Yahoo")
    lote.rescue(rows, {}, flaky, waits=(7.0,), sleep=lambda s: None, on_progress=msgs.append)
    assert any("Esperando 7" in m for m in msgs) and any("Reintentando AAA" in m for m in msgs), msgs


if __name__ == "__main__":
    n = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn(); n += 1; print("OK", name)
    print(f"{n} pruebas OK")


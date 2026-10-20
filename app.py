"""
DCF v2 — Valuación por escenarios (Pesimista / Base / Optimista) para cualquier acción listada en Yahoo Finance.

Modelos:  FCFF desde ingresos (no financieras) · Retorno excedente (bancos, seguros, mREITs, utilities reguladas) · AFFO (REITs).
Ejecutar:  streamlit run app.py
"""

import io
from datetime import date, datetime

import altair as alt
import pandas as pd
import streamlit as st

try:
    import data_yahoo
    from country_risk import country_crp
    from country_tax import marginal_tax_for
    from industry_data import (IND, ERP_DAMODARAN, ERP_DATE, ERP_MAX_MESES, erp_is_stale, DATE as IND_DATE, DATE_ISO as IND_DATE_ISO,
                               attach as attach_industry, get_ref, industry_names, reference_status)
    from dcf_models import (MOAT_SPREADS, BETA_MIN, BETA_MAX, ZONE_NAMES, price_zones, implied_wacc, field_checks, margin_options, Params, data_checks, route_model,
                            run_scenarios, sensitivity, terminal_growth_default)
    from methods_map import methods_map as build_methods_map, DORSEY_G_DEFAULT, DORSEY_G_RANGE, DORSEY_R_MIN
    from industry_multiples import MULT_DATE
except ImportError as _imp_err:   # archivos de versiones distintas en GitHub, o módulo viejo en memoria: diagnóstico claro
    import os as _os
    _aqui = _os.path.dirname(_os.path.abspath(__file__))
    _marcas = {"dcf_models.py": "clean_margin_trend", "industry_data.py": "USE_DAMODARAN_LIST", "damodaran_map.py": "def lookup_country", "data_yahoo.py": "profile_fallback", "profile_fallback.py": "def country_chain", "lote.py": "methods_map",
               "methods_map.py": "def methods_map", "industry_multiples.py": "def get_multiples"}

    def _tiene(archivo, marca):
        try:
            with open(_os.path.join(_aqui, archivo), encoding="utf-8") as _f:
                return marca in _f.read()
        except OSError:
            return None
    _estado = {k: _tiene(k, m) for k, m in _marcas.items()}
    _viejos = [k for k, v in _estado.items() if v is False]
    _faltan = [k for k, v in _estado.items() if v is None]
    if _viejos or _faltan:
        _dx = ("**Archivos en el repositorio que NO tienen la versión nueva:** " + ", ".join(_viejos) + ". " if _viejos else "") + \
              ("**No encontré:** " + ", ".join(_faltan) + ". " if _faltan else "") + \
              "Subí esos archivos con su nombre exacto (sin «(1)») a la misma carpeta que app.py."
    else:
        _dx = ("**Los archivos del repositorio SÍ son los nuevos.** La app tiene guardada en memoria una versión vieja: "
               "abrí «Manage app» (abajo a la derecha) → menú ⋮ → **Reboot app**.")
    st.error(f"**Versiones distintas o módulo viejo en memoria.** {_dx}\n\nDetalle técnico: {_imp_err}")
    st.stop()


if not getattr(data_yahoo, "PROFILE_FALLBACK", False):
    st.error("**data_yahoo.py del repositorio es una versión vieja** (no tiene los respaldos de país de la v2.50). "
             "Subilo de nuevo desde el zip v2.50, junto con **profile_fallback.py** (archivo nuevo), y hacé Reboot app.")
    st.stop()


def _stretch():
    """Compatibilidad entre versiones de Streamlit (use_container_width fue reemplazado por width='stretch')."""
    try:
        major, minor = (int(x) for x in st.__version__.split(".")[:2])
        if (major, minor) >= (1, 50):
            return {"width": "stretch"}
    except Exception:  # noqa: BLE001
        pass
    return {"use_container_width": True}


STRETCH = _stretch()

# Código de colores: azul = dato importado · amarillo = supuesto (editable o estimación de la app) · verde = cálculo
KIND_STYLE = {"dato": "background-color:#dbe9f6;color:#1b2a3a",
              "supuesto": "background-color:#fff4c2;color:#3a3200",
              "calc": "background-color:#e1f2e1;color:#1d3320"}
_DATO = ("Margen EBIT actual", "Crecimiento ingresos año", "Caja (M)", "Deuda (M)", "ROE TTM", "Book value", "AFFO por acción")
_SUPUESTO = ("g terminal", "Margen EBIT objetivo", "Foso", "Persistencia", "Retorno de lo reinvertido")


def kind_of(label: str) -> str:
    if any(label.startswith(k) for k in _DATO):
        return "dato"
    if any(label.startswith(k) for k in _SUPUESTO):
        return "supuesto"
    return "calc"


def color_rows(df, kind_fn=kind_of):
    """Pinta cada fila según su tipo (el índice es la etiqueta de la fila)."""
    return df.style.apply(lambda r: [KIND_STYLE[kind_fn(str(r.name))]] * len(r), axis=1)


COLOR_LEGEND = ("🟦 **Dato importado** · 🟨 **Supuesto** (editable en la barra lateral o estimación de la app) · "
                "🟩 **Cálculo**")
STATE_ICON = {"ok": "✅", "aviso": "⚠️", "info": "ℹ️"}

APP_VERSION = "v2.53 · 20/10/2026"

st.set_page_config(page_title="DCF v2 · Valuación por escenarios", page_icon="📈", layout="wide")
if not (hasattr(data_yahoo, "statement_currency_check") and hasattr(data_yahoo, "currency_diagnostic")):
    st.warning("**data_yahoo.py es una versión vieja:** falta el control o el diagnóstico de la moneda de los estados contables (v2.47.1). "
               "Subí data_yahoo.py de la v2.47 junto con los demás archivos y reiniciá la app.")

MODEL_LABELS = {"FCFF": "Empresa común (todos los demás sectores)",
                "ER": "Banco / aseguradora / financiera / utility regulada",
                "REIT": "REIT (inmobiliaria que cobra alquileres)"}
SCEN_ORDER = ["Pesimista", "Base", "Optimista"]


# ----------------------------------------------------------------------------
class _Degraded(Exception):
    """Carga incompleta: se devuelve igual, pero NO se guarda en la caché (así volver a intentar reintenta de verdad)."""
    def __init__(self, payload):
        super().__init__("carga incompleta")
        self.payload = payload


@st.cache_data(ttl=3600, show_spinner=False)
def _cached_load_ok(sym: str, allow_annual: bool, adr_ratio: float, fin_mode: str = "auto"):
    if fin_mode == "auto":
        co_, warns_ = data_yahoo.load_company(sym, allow_annual, adr_ratio or None)          # compatible con un data_yahoo.py viejo
    else:
        try:
            co_, warns_ = data_yahoo.load_company(sym, allow_annual, adr_ratio or None, fin_mode)
        except TypeError as _e:
            if "positional argument" in str(_e):
                raise RuntimeError("data_yahoo.py es una versión vieja (no conoce la opción de moneda): subí data_yahoo.py de la v2.47 "
                                   "a la misma carpeta que app.py y reiniciá la app.") from _e
            raise
    if data_yahoo.is_degraded(co_):
        raise _Degraded((co_, warns_))
    return co_, warns_


def cached_load(sym: str, allow_annual: bool, adr_ratio: float, fin_mode: str = "auto"):
    try:
        return _cached_load_ok(sym, allow_annual, adr_ratio, fin_mode)
    except _Degraded as e:
        return e.payload


def pct(x, d=1):
    return "n/d" if x is None or x != x else f"{x * 100:.{d}f}%"


def money(x, cur="", d=2):
    return "n/d" if x is None or x != x else f"{x:,.{d}f} {cur}".strip()


PRICE_COLOR = "#4fc3f7"   # celeste del marcador de precio
ZONE_COLORS = ["#1a7f37", "#7cc242", "#f2c230", "#f28c28", "#d63a2f"]
ZONE_PHRASES = [
    "🟢 **Muy infravalorada:** el precio está por debajo incluso del escenario pesimista.",
    "🟢 **Infravalorada:** el precio está entre el escenario pesimista y el punto medio hacia el base.",
    "🟡 **Bien valuada:** el precio está cerca del caso base, dentro del rango razonable de los escenarios.",
    "🟠 **Sobrevalorada:** el precio está entre el punto medio hacia el optimista y el escenario optimista: el mercado ya descuenta casi lo mejor.",
    "🔴 **Muy sobrevalorada:** el precio está por encima del escenario optimista: el mercado paga más que el mejor caso del modelo.",
]


def zones_of(price, res):
    vp, vb, vo = (res[n].per_share if res[n].ok else None for n in SCEN_ORDER)
    return price_zones(vp, vb, vo, price)


def verdict(price, res):
    """Frase simple: en qué zona está el precio respecto de los tres escenarios."""
    z = zones_of(price, res)
    if z is None:
        return "No hay los tres escenarios disponibles para ubicar el precio."
    return ZONE_PHRASES[z["zone"]]


def zone_bar_html(z, price, cur=""):
    """Barra de 5 colores con el precio actual como marcador celeste bien visible (HTML simple, sin librerías)."""
    cuts, base = z["cuts"], z["base"]
    lo = min(cuts[0], price) * 0.8
    hi = max(cuts[3], price) * 1.1
    if hi <= lo:
        hi = lo + 1.0
    span = hi - lo
    edges = [lo] + cuts + [hi]
    pos = lambda v: max(0.0, min(100.0, (v - lo) / span * 100))
    segs = "".join(f'<div style="width:{(edges[k + 1] - edges[k]) / span * 100:.3f}%;background:{ZONE_COLORS[k]}" '
                   f'title="{ZONE_NAMES[k]}"></div>' for k in range(5))

    def tick(v, name):
        return (f'<div style="position:absolute;left:{pos(v):.2f}%;top:34px;transform:translateX(-50%);text-align:center;'
                f'font-size:13px;line-height:1.3;white-space:nowrap">{name}<br><b style="font-size:15px">{v:,.2f}</b></div>')
    px = pos(price)
    marker = (f'<div style="position:absolute;left:{px:.2f}%;top:-8px;width:5px;height:42px;transform:translateX(-50%);'
              f'background:{PRICE_COLOR};border:1px solid #fff;border-radius:3px;box-shadow:0 0 0 1px rgba(0,0,0,.35)"></div>'
              f'<div style="position:absolute;left:{px:.2f}%;top:-46px;transform:translateX(-50%);background:{PRICE_COLOR};color:#04283a;'
              f'font-size:14px;font-weight:700;line-height:1;padding:5px 9px;border-radius:12px;white-space:nowrap;'
              f'box-shadow:0 1px 3px rgba(0,0,0,.35)">Precio {price:,.2f} {cur}</div>')
    if cuts[3] <= 0:        # los tres escenarios valen 0: una sola etiqueta en vez de tres encimadas
        ticks = ('<div style="position:absolute;left:0;top:34px;font-size:13px;white-space:nowrap">'
                 'Pesimista, base y optimista: <b>0</b> por acción</div>')
    else:
        ticks = tick(cuts[0], "Pesimista") + tick(base, "Base") + tick(cuts[3], "Optimista")
    return ('<div style="position:relative;margin:56px 14px 66px 14px;">'
            f'<div style="display:flex;height:22px;border-radius:11px;overflow:hidden;border:1px solid rgba(128,128,128,.4)">{segs}</div>'
            + marker + ticks + '</div>')


def build_excel(co, res, base, sens, p, model):
    """Excel con resultados, supuestos, proyección base y sensibilidad (para guardar o comparar empresas)."""
    buf = io.BytesIO()
    resumen = pd.DataFrame([{"Escenario": n, "Valor por acción": res[n].per_share if res[n].ok else None,
                             "vs precio": (res[n].per_share / co.price - 1) if res[n].ok else None,
                             "Tasa de descuento": res[n].discount_rate if res[n].ok else None,
                             "g terminal": res[n].g_terminal if res[n].ok else None} for n in SCEN_ORDER])
    params = pd.DataFrame([
        ("Ticker", co.ticker), ("Empresa", co.name), ("Sector", co.sector), ("Industria", co.industry),
        ("Modelo", model), ("Precio", co.price), ("Moneda", co.currency), ("Cap. bursátil (M)", co.market_cap_m),
        ("Tasa libre de riesgo", p.rf if p.rf is not None else co.rf), ("Prima de mercado", p.erp),
        ("Prima de riesgo país", p.crp), ("Tasa marginal de impuestos", p.marginal_tax),
        ("Foso (moat)", base.extras.get("moat_label", p.moat) if base.extras else p.moat), ("g terminal", p.g_terminal),
        ("Horizonte (años)", p.horizon),
        ("Datos al", co.asof), ("Fuente TTM", co.ttm_source), ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ], columns=["Parámetro", "Valor"])
    controles = pd.DataFrame([{"Estado": {"ok": "OK", "aviso": "REVISAR", "info": "Dato faltante"}[c["Estado"]],
                               "Dato": c["Campo"], "Valor usado": c["Valor"], "Qué controlar": c["Control"]}
                              for c in field_checks(co, model, p)])
    leyenda = pd.DataFrame([("Azul", "Dato importado"), ("Amarillo", "Supuesto (editable o estimación de la app)"),
                            ("Verde", "Cálculo")], columns=["Color", "Significado"])
    _supuestos = ("Prima de mercado", "Foso (moat)", "g terminal", "Horizonte (años)")
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        resumen.to_excel(xw, sheet_name="Resumen", index=False)
        params.to_excel(xw, sheet_name="Parametros", index=False)
        pd.DataFrame(base.table).to_excel(xw, sheet_name="Proyeccion_Base", index=False)
        if sens.get("ok"):
            pd.DataFrame(sens["grid"], index=[f"{x:.2%}" for x in sens["rates"]],
                         columns=[f"g {x:.2%}" for x in sens["gs"]]).to_excel(xw, sheet_name="Sensibilidad")
        controles.to_excel(xw, sheet_name="Controles_de_datos", index=False)
        leyenda.to_excel(xw, sheet_name="Leyenda_de_colores", index=False)
        try:
            from openpyxl.styles import PatternFill
            fills = {"dato": PatternFill("solid", fgColor="DBE9F6"), "supuesto": PatternFill("solid", fgColor="FFF4C2"),
                     "calc": PatternFill("solid", fgColor="E1F2E1")}
            wsp = xw.book["Parametros"]
            for r in range(2, wsp.max_row + 1):
                kind = "supuesto" if wsp.cell(r, 1).value in _supuestos else "dato"
                for c in (1, 2):
                    wsp.cell(r, c).fill = fills[kind]
            for name in ("Resumen", "Proyeccion_Base"):
                for c in xw.book[name][1]:
                    c.fill = fills["calc"]
            wsl = xw.book["Leyenda_de_colores"]
            for r, k in ((2, "dato"), (3, "supuesto"), (4, "calc")):
                wsl.cell(r, 1).fill = fills[k]
        except Exception:  # noqa: BLE001
            pass
    return buf.getvalue()


@st.cache_data(ttl=86400, show_spinner=False)
def get_logo(sym: str, website: str):
    """Busca el logo de la empresa (solo decoración). Devuelve bytes de imagen o None; nunca falla."""
    try:
        import requests
        from urllib.parse import urlparse
        host = urlparse(website if "//" in (website or "") else "//" + (website or "")).netloc.replace("www.", "")
        base = sym.split(".")[0].upper()
        cands = [f"https://financialmodelingprep.com/image-stock/{base}.png"]
        if host:
            cands += [f"https://www.google.com/s2/favicons?domain={host}&sz=128",
                      f"https://icons.duckduckgo.com/ip3/{host}.ico"]
        for u in cands:
            try:
                r = requests.get(u, timeout=4, headers={"User-Agent": "Mozilla/5.0"})
                if r.status_code == 200 and r.headers.get("content-type", "").startswith("image") and len(r.content) > 1200:
                    return r.content
            except Exception:  # noqa: BLE001
                continue
    except Exception:  # noqa: BLE001
        pass
    return None


# ----------------------------------------------------------------------------
# Barra lateral: supuestos globales
# ----------------------------------------------------------------------------
def K(name):
    """Clave de widget que cambia con cada ticker nuevo => los valores se reinician."""
    return f"{name}_{st.session_state.get('rn', 0)}"


_mode = st.sidebar.radio("Modo", ["Una empresa", "Lote de empresas"], key="app_mode",
                         help="«Lote de empresas» valúa muchas empresas de una vez y devuelve un solo Excel (para calibrar el modelo).")
if _mode == "Lote de empresas":
    import lote
    lote.render(cached_load, STRETCH, APP_VERSION)
    st.stop()

with st.sidebar:
    st.header("Supuestos")
    horizon = st.slider("Horizonte explícito (años)", 5, 15, 10, key=K("horizon"))
    g_term_in = st.number_input("Crecimiento terminal g (%) — 0 = automático", 0.0, 4.5, 0.0, 0.25, key=K("gterm"),
                                help="Crecimiento perpetuo después del último año proyectado. Automático: tasa libre de riesgo − 1 punto, "
                                     "entre 3,0 % y 3,5 % (techo por debajo del PIB nominal de EE.UU., ~3,8–4,1 %: una empresa sola no crece para siempre "
                                     "como toda la economía). Para flujos en euros (~3 %), yenes o francos suizos (~1,5–2,5 %) conviene "
                                     "cargar uno menor a mano.") / 100
    erp = st.number_input("Prima de riesgo de mercado (%)", 3.0, 8.0, round(ERP_DAMODARAN * 100, 2), 0.05, key=K("erp"),
                          format="%.2f",
                          help=f"Por defecto: la prima implícita de EE.UU. de Damodaran ({ERP_DAMODARAN:.2%}, {ERP_DATE}), medida sobre el bono "
                               f"a 10 años. Se actualiza a mano cada 4-6 meses; la app avisa pasados {ERP_MAX_MESES} meses.") / 100
    crp_auto = st.checkbox("Prima de riesgo país automática", value=True, key=K("crpauto"),
                           help="Según el país de la empresa, con la tabla de Damodaran (enero 2026). "
                                "EE.UU., Europa desarrollada y Japón dan casi 0; emergentes suman varios puntos.")
    crp_manual = 0.0
    if not crp_auto:
        crp_manual = st.number_input("Prima de riesgo país (%)", 0.0, 30.0, 0.0, 0.25, key=K("crpman")) / 100
    moat = st.selectbox("Ventaja competitiva (moat)", ["Automático"] + list(MOAT_SPREADS), key=K("moat"),
                        help="Cuánto por encima del costo del capital puede seguir ganando la empresa sobre el capital que reinvierte, "
                             "para siempre. Automático: la app lo estima con tres criterios (exceso de retorno sobre el costo del capital, persistencia del margen "
                             "y comparación con sus pares); abajo del resultado ves las razones. Podés forzar Ninguno / Estrecho / Amplio. "
                             "Es el supuesto que más mueve el valor.")
    allow_annual = st.checkbox("Permitir último año fiscal si no hay TTM", value=False, key=K("annual"),
                               help="Por defecto se exigen datos de los últimos 12 meses (TTM).")
    adr_ratio = st.number_input("Acciones ordinarias por ADR (0 = no es ADR / automático)", 0.0, 1000.0, 0.0, 1.0, key=K("adr"),
                                help="SOLO para ADRs de empresas extranjeras. Dejalo en 0 para empresas de EE.UU. "
                                     "(KO, JPM, V, O…). Ej.: TSM = 5, ITUB = 1, BABA = 8. Se ignora si la empresa "
                                     "reporta en dólares.")
    _fin_opts = {"Automática (con control contra Yahoo)": "auto",
                 "Ya están en la moneda de cotización (no convertir)": "no_convertir",
                 "Convertir siempre (confiar en la etiqueta de Yahoo)": "convertir"}
    fin_mode = _fin_opts[st.selectbox("Moneda de los estados contables", list(_fin_opts), key=K("fincur"),
                                      help="Para empresas que reportan en otra moneda que la de cotización (ADRs). Yahoo a veces rotula mal la moneda "
                                           "(ej.: Vale y Petrobras dicen BRL pero los montos ya están en dólares) y los valores quedaban 5 veces chicos. "
                                           "«Automática» compara precio/ventas, precio/valor libro y precio/ganancia con los de Yahoo (vota cada señal) y corrige sola cuando al menos dos coinciden con «no convertir» y ninguna con «convertir»; avisa cuando lo hace.")]
    with st.expander("Ajustes avanzados"):
        margin_mode = st.selectbox("Margen operativo objetivo", ["median", "ttm", "analysts", "manual"], key=K("mmode"),
                                   format_func={"median": "Mediana histórica (3-4 años + TTM)",
                                                "ttm": "Mantener TTM",
                                                "analysts": "Según analistas (expansión de margen esperada)",
                                                "manual": "Manual"}.get,
                                   help="Si no hay estimaciones de analistas, «Según analistas» usa la mediana histórica.")
        margin_manual = None
        if margin_mode == "manual":
            margin_manual = st.number_input("Margen EBIT objetivo (%)", -20.0, 80.0, 15.0, 0.1, format="%.1f", key=K("mman")) / 100
        margin_years = st.slider("Años para converger al margen objetivo", 1, 10, 5, key=K("myears"))
        use_g = st.checkbox("Fijar crecimiento de ingresos años 1 y 2", key=K("usg"))
        growth_override = None
        if use_g:
            g1 = st.number_input("Crecimiento año 1 (%)", -20.0, 50.0, 8.0, 0.1, format="%.1f", key=K("g1")) / 100
            g2 = st.number_input("Crecimiento año 2 (%)", -20.0, 50.0, 6.0, 0.1, format="%.1f", key=K("g2")) / 100
            growth_override = (g1, g2)
        rf_in = st.number_input("Tasa libre de riesgo 10 años (%) — 0 = automática", 0.0, 15.0, 0.0, 0.05,
                                key=K("rf"), help="Dejala en 0: la app la busca sola (Yahoo / Reserva Federal).")
        beta_manual = st.number_input("Beta — 0 = la de Yahoo", 0.0, 3.0, 0.0, 0.05, key=K("beta"))
        wacc_manual = st.number_input("WACC fijo (%) — 0 = calculado", 0.0, 20.0, 0.0, 0.25, key=K("wacc")) / 100
        persistence = st.slider("Persistencia del exceso de ROE (financieras)", 0.0, 1.0, 0.5, 0.05, key=K("pers"),
                                help="Qué fracción del ROE por encima de Ke sobrevive en el largo plazo.")
        affo_ratio = st.slider("AFFO / FFO estimado (REITs)", 0.6, 1.0, 0.90, 0.01, key=K("affor"))
        affo_ps = st.number_input("AFFO por acción reportado — 0 = estimar", 0.0, 1000.0, 0.0, 0.05, key=K("affops"))
        incl_nonop = st.checkbox("Sumar inversiones en otras empresas (valor contable)", value=True, key=K("nonop"),
                                 help="Participaciones y valores de largo plazo que no generan ventas (ej.: las embotelladoras de "
                                      "KO, las inversiones de Amazon). Como el resultado operativo no las incluye, se suman aparte.")
        use_ind_beta = st.checkbox("Usar la beta de la industria (Damodaran)", value=True, key=K("indbeta"),
                                   help="La beta usada es el promedio entre la del sector (ajustada por la deuda de esta empresa) y la propia "
                                        f"de la acción, con piso {BETA_MIN:.1f} y tope {BETA_MAX:.1f}. La del sector es mucho más estable que la de una sola acción "
                                        "(ej.: 0,29 en KO). Si la desmarcás o no hay datos del sector, se usa la de la acción con ajuste de Blume.")
        cap_rd = st.checkbox("Capitalizar el gasto en I+D (Damodaran)", value=True, key=K("caprd"),
                             help="Trata el I+D como inversión (activo que se amortiza) y no como gasto: sube el resultado operativo y el capital "
                                  "invertido de farmacéuticas y tecnológicas. Usa los años de historia que entrega Yahoo (4). Sin dato de I+D no hace nada.")
        surv_adj = st.checkbox("Ajustar por riesgo de quiebra (Damodaran)", value=True, key=K("surv"),
                               help="Si la calificación sintética es peor que BBB, el valor se multiplica por (1 − probabilidad de quiebra a 10 años), "
                                    "con la tabla de Damodaran. En una quiebra se supone que el accionista no recibe nada.")
        ind_manual = st.selectbox("Industria para comparar (Damodaran)", ["Automática"] + industry_names(),
                                  key=K("dind"),
                                  help="Automática: la app busca la empresa en la planilla de Damodaran (su industria exacta) y, si no figura, usa la industria de Yahoo. Elegila acá si no la detecta o si querés compararla con otra.")
        lease_opt = st.selectbox("Arrendamientos como deuda", ["auto", "sí", "no"], key=K("lease"),
                                 help="Auto: EE.UU. = no (el gasto de alquiler ya está en el EBIT); resto = sí.")

try:
    _rs = reference_status()
    _venc = [r for r in _rs if r["vencida"]]
    with st.sidebar.expander(("⚠️ Datos de referencia: hay que renovar " + str(len(_venc))) if _venc else "✅ Datos de referencia al día"):
        for r in _rs:
            st.markdown(f"{'⚠️' if r['vencida'] else '✅'} **{r['nombre']}** · {r['fecha']} · {r['meses']:.0f} meses "
                        f"(se renueva cada {r['limite_meses']})" + (f"  \n_{r['donde']}_" if r["vencida"] else ""))
except Exception:  # noqa: BLE001
    pass
st.sidebar.caption(f"Versión de la app: {APP_VERSION}")
st.title("📈 DCF v2 — Valuación por escenarios")
st.caption("Ingresá un ticker de Yahoo Finance (ej.: AAPL, KO, JPM, O, TSM, SAP.DE, 7203.T, NESN.SW, ITUB). "
           "Resultado: valor por acción en tres escenarios. Herramienta educativa, no es una recomendación.")

with st.form("f"):
    c1, c2 = st.columns([3, 3])
    ticker = c1.text_input("Ticker", value=st.session_state.get("last", "")).strip().upper()
    model_choice = c2.selectbox("Modelo", ["auto", "FCFF", "ER", "REIT"], key=K("model"),
                                format_func=lambda k: "Automático (detecta el sector)" if k == "auto" else MODEL_LABELS[k],
                                help="Dejalo en Automático: la app detecta el sector y elige el modelo. "
                                     "Cambialo solo si el aviso te dice que no pudo detectarlo.")
    go = st.form_submit_button("Valuar", type="primary")

st.caption("📌 La tasa libre de riesgo (bono del Tesoro de EE.UU. a 10 años) se busca sola. Si querés cargarla vos, "
           "está en la barra lateral → Ajustes avanzados. Dónde mirarla: "
           "[FRED – Reserva Federal](https://fred.stlouisfed.org/series/DGS10) · "
           "[CNBC – en vivo](https://www.cnbc.com/quotes/US10Y).")

if go and ticker:
    if ticker != st.session_state.get("last", ""):
        # Ticker nuevo: la barra lateral y el modelo vuelven a sus valores por defecto
        st.session_state["rn"] = st.session_state.get("rn", 0) + 1
        st.session_state["last"] = ticker
        st.rerun()
    st.session_state["last"] = ticker
sym = st.session_state.get("last", "")
if not sym:
    st.info("Escribí un ticker y apretá **Valuar**.")
    st.stop()

# ----------------------------------------------------------------------------
# Datos
# ----------------------------------------------------------------------------
try:
    with st.spinner(f"Descargando datos de {sym}…"):
        co, load_warns = cached_load(sym, allow_annual, adr_ratio, fin_mode)
except Exception as e:  # noqa: BLE001
    st.error(f"No pude cargar **{sym}**: {e}")
    st.stop()

_CAP = "capitalización de mercado (posible ADR sin ajustar)"
_miss = data_yahoo.missing_fields(co)
if _miss:
    _effects = {
        "industria": "sin industria no hay beta ni referencias del sector (el WACC y el foso quedan menos confiables)",
        "país": "sin país se usa un impuesto marginal de 25 % (supuesto), prima de riesgo país 0 y los arrendamientos se cuentan como deuda",
        "sector": "sin sector se pierde el aviso de industrias cíclicas",
        "nombre": "sin nombre solo se muestra el ticker",
        _CAP: "las acciones son las ORDINARIAS y no las del ADR que cotiza: el valor por acción puede estar mal por un múltiplo (ej.: TSM = 5, TM = 10). "
              "No uses este valor; reintentá, o cargá 'ordinarias por ADR' en la barra lateral",
    }
    _list = _miss[0] if len(_miss) == 1 else ", ".join(_miss[:-1]) + " y " + _miss[-1]
    _detail = "; ".join(_effects[m] for m in _miss) + "."
    _retry = " Probá de nuevo en unos minutos: esta carga no se guarda, así que el reintento vuelve a consultar. Detalle en los avisos de más abajo."
    if "industria" in _miss or _CAP in _miss:
        st.error(f"🔴 **Resultado degradado: no pude leer {_list} de la empresa en Yahoo.** {_detail[0].upper() + _detail[1:]}{_retry}")
    else:
        _tip = (" Si la empresa no es de EE.UU., cargá la prima de riesgo país en la barra lateral." if "país" in _miss else "")
        st.warning(f"🟡 **Dato incompleto: no pude leer {_list} de la empresa en Yahoo.** "
                   f"{_detail[0].upper() + _detail[1:]} El efecto en la valuación es acotado.{_tip}{_retry}")

if getattr(co, "country_from", ""):
    st.info(f"ℹ️ Yahoo no entregó el país: se tomó **{co.country}** de {co.country_from}, para no perder la prima de riesgo país y el impuesto del país.")
    if co.rev_g_y0 is None and co.rev_g_y1 is None:
        st.warning("🟡 Yahoo limitó varios módulos a la vez (tampoco llegaron las estimaciones de analistas): el valor puede cambiar al reintentar. "
                   "Esta carga no se guarda; probá de nuevo en unos minutos.")

if hasattr(data_yahoo, "currency_diagnostic") and (co.fin_currency != co.currency or any("Moneda de los estados contables" in w for w in load_warns)):
    with st.expander("🔎 Diagnóstico de la moneda de los estados contables (compara Yahoo con la app)"):
        st.caption("Solo consulta, no cambia ningún valor. Sirve para revisar si Yahoo rotuló bien la moneda de los estados contables.")
        if st.button("Consultar a Yahoo", key=K("diagmon")):
            try:
                _dg = data_yahoo.currency_diagnostic(co.ticker)
                _mc, _rev = co.market_cap_m, co.revenue_ttm
                _fxu = co.fx if co.fx else 1.0
                _rows = [("Moneda de cotización", co.currency), ("Moneda de los estados contables que usó la app", co.fin_currency),
                         ("Tipo de cambio aplicado a los montos", f"{co.fx:.4f}"),
                         ("Ingresos TTM de la app (M, ya en moneda de cotización)", f"{_rev:,.0f}"),
                         ("Capitalización de la app (M)", f"{_mc:,.0f}"),
                         ("Precio/ventas de la app", f"{(_mc / _rev):.2f}" if _rev else "-"),
                         ("Precio/ventas de la app SI los montos no se hubieran convertido", f"{(_mc / (_rev / _fxu)):.2f}" if _rev else "-"),
                         ("Patrimonio de la app (M)", f"{co.equity:,.0f}" if co.equity else "-"),
                         ("Precio/valor libro de la app", f"{(_mc / co.equity):.2f}" if co.equity else "-"),
                         ("Ganancia neta TTM de la app (M)", f"{co.ni_ttm:,.0f}" if co.ni_ttm else "-"),
                         ("Precio/ganancia de la app", f"{(_mc / co.ni_ttm):.1f}" if co.ni_ttm else "-")]
                _rows += [(k, ("-" if v is None else (f"{v:,.2f}" if isinstance(v, float) else str(v)))) for k, v in _dg.items()]
                st.table(pd.DataFrame(_rows, columns=["Dato", "Valor"]).set_index("Dato"))
            except Exception as _e:  # noqa: BLE001
                st.error(f"No pude consultar a Yahoo: {_e}")

if crp_auto:
    crp, crp_note = country_crp(co.country)
else:
    crp, crp_note = crp_manual, "Cargada a mano."
auto_model, why = route_model(co)
model = auto_model if model_choice == "auto" else model_choice

tax_marg, tax_note = marginal_tax_for(co.country)
_rf_pre = (rf_in / 100) or co.rf
g_auto = terminal_growth_default(_rf_pre)
g_used = g_term_in if g_term_in > 0 else g_auto

p = Params(
    horizon=horizon, g_terminal=g_used, erp=erp, crp=crp, moat="auto" if moat == "Automático" else moat,
    marginal_tax=tax_marg,
    margin_mode=margin_mode, margin_manual=margin_manual, margin_years=margin_years,
    growth_override=growth_override, rf=(rf_in / 100) or None, beta=beta_manual or None,
    wacc_override=wacc_manual or None, persistence=persistence, affo_ratio=affo_ratio,
    affo_ps_override=affo_ps or None,
    include_leases={"auto": None, "sí": True, "no": False}[lease_opt],
    use_industry_beta=use_ind_beta, include_nonop=incl_nonop, capitalize_rd=cap_rd, survival_adj=surv_adj,
)

_logo = get_logo(co.ticker, co.website)
if _logo:
    _c1, _c2 = st.columns([1, 9], vertical_alignment="center")
    _c1.image(_logo, width=64)
    _c2.subheader(f"{co.name} ({co.ticker})")
else:
    st.subheader(f"{co.name} ({co.ticker})")
st.write(f"**Sector:** {co.sector or 'n/d'} · **Industria:** {co.industry or 'n/d'} · **País:** {co.country or 'n/d'} · "
         f"**Precio:** {money(co.price, co.currency)} · **Cap. bursátil:** {money(co.market_cap_m, co.currency, 0)} M · "
         f"**Datos contables al:** {co.asof or 'n/d'}")
_rf_used = p.rf if p.rf is not None else co.rf
st.write(f"**Tasa libre de riesgo usada:** {_rf_used:.2%} "
         f"({'cargada por vos' if p.rf is not None else 'automática'})")
st.write(f"**Prima de riesgo país aplicada:** {crp:.2%} — {crp_note}")
st.write(f"**Tasa marginal de impuestos:** {tax_marg:.2%} — {tax_note}")
g_slot = st.empty()                                   # se completa después de valuar: muestra el g realmente aplicado
if model_choice == "auto":
    st.info(f"**Modelo: {MODEL_LABELS[model]}.** {why}")
elif model != auto_model:
    st.warning(f"Modelo elegido a mano: {MODEL_LABELS[model]}. El automático sería {MODEL_LABELS[auto_model]}. {why}")

co0 = co
co, ind_name, ind_how = attach_industry(co0, model, None if ind_manual == "Automática" else ind_manual)
res = run_scenarios(co, p, model)
base = res["Base"]
if (model_choice == "auto" and model == "FCFF" and not base.ok and "EBIT" in base.error
        and (co.ni_ttm or 0) > 0 and (co.equity or 0) > 0):
    model = "ER"
    why = ("Sin resultado operativo pero con ganancias y patrimonio (típico de bancos y aseguradoras): "
           "se pasó automáticamente a retorno excedente. Confirmalo si no es una financiera.")
    co, ind_name, ind_how = attach_industry(co0, model, None if ind_manual == "Automática" else ind_manual)
    res = run_scenarios(co, p, model)
    base = res["Base"]

_g_txt = "cargado por vos" if g_term_in > 0 else "automático: tasa libre de riesgo − 1 punto, entre 3,0 % y 3,5 %"
if base.ok and abs(base.g_terminal - g_used) > 5e-5:
    g_slot.write(f"**Crecimiento terminal:** {g_used:.2%} pedido ({_g_txt}) → **{base.g_terminal:.2%} aplicado**: "
                 "el modelo mantiene una distancia mínima con la tasa de descuento (ver avisos).")
else:
    g_slot.write(f"**Crecimiento terminal:** {g_used:.2%} ({_g_txt})")

all_warns = list(dict.fromkeys(load_warns + data_checks(co)))
if model == "FCFF" and (co.sector or "") in ("Energy", "Basic Materials"):
    all_warns.append("Sector cíclico: el margen EBIT está normalizado con la mediana histórica, pero revisá "
                     "si el ciclo actual es pico o piso.")
if crp_auto and (not co.country or "no figura" in crp_note):
    all_warns.append(crp_note)

if g_used > 0.03 and (co.currency or "").upper() in ("EUR", "JPY", "CHF", "SEK", "DKK", "NOK") :
    all_warns.append(f"Crecimiento terminal de {g_used:.2%} con cotización en {co.currency}: el crecimiento nominal de largo plazo "
                     "de esa economía es menor (≈1,5–3 %). Considerá cargar un g menor en la barra lateral.")
try:
    if erp_is_stale():
        all_warns.append(f"La prima de riesgo de mercado cargada ({ERP_DAMODARAN:.2%}) es de {ERP_DATE} y tiene más de {ERP_MAX_MESES} meses: "
                         "conviene actualizarla con la prima implícita vigente de Damodaran (ERPbymonth, columna ERP T12m).")
except Exception:  # noqa: BLE001
    pass
try:
    _dam = [r for r in reference_status() if r["vencida"] and r["nombre"].startswith("Planilla")]
    if _dam:
        all_warns.append(f"La planilla de empresas por industria de Damodaran es de {_dam[0]['fecha']} y tiene más de 12 meses: "
                         "hay empresas nuevas o cambiadas de industria que no figuran. Conviene bajarla de nuevo.")
except Exception:  # noqa: BLE001
    pass
try:
    if (date.today() - date.fromisoformat(IND_DATE_ISO)).days > 365:
        all_warns.append(f"Las tablas de Damodaran (prima de riesgo país, industrias, impuestos) son de {IND_DATE} y tienen más "
                         "de 12 meses: conviene actualizarlas.")
except Exception:  # noqa: BLE001
    pass

_n_check = sum(1 for c in field_checks(co, model, p) if c["Estado"] == "aviso")
if _n_check:
    all_warns.append(f"Controles por dato: {_n_check} para revisar (pestaña «Controles de cordura»).")

if not base.ok:
    st.error(f"**No se puede valuar con este modelo:** {base.error}")
    for w in all_warns:
        st.warning(w)
    st.stop()

# ----------------------------------------------------------------------------
# Resultado principal
# ----------------------------------------------------------------------------
cols = st.columns(3)
for col, name in zip(cols, SCEN_ORDER):
    r = res[name]
    with col:
        if r.ok:
            up = r.per_share / co.price - 1
            st.metric(f"{name}", money(r.per_share, co.currency), f"{up:+.0%} vs precio",
                      delta_color="normal")
        else:
            st.metric(f"{name}", "n/d")
            st.caption(r.error)

_lc = base.extras.get("lifecycle")
if _lc:
    st.caption(f"🧭 **Etapa: {_lc}.** Según Damodaran, lo que más pesa en su valor: {base.extras.get('lifecycle_factors', '')}.")
if base.extras.get("p_default"):
    st.warning(f"⚠️ **Riesgo de quiebra:** calificación sintética {base.extras['rating']}; el valor ya está multiplicado por "
               f"(1 − {base.extras['p_default']:.0%}). Como negocio en marcha valdría {money(base.extras['per_share_going'], co.currency)} por acción (Base).")

_z = zones_of(co.price, res)
if _z is not None:
    st.markdown(zone_bar_html(_z, co.price, co.currency), unsafe_allow_html=True)
    st.caption("Zonas calculadas con los tres escenarios de esta empresa (pesimista, base y optimista): cuanto más incierta, más ancha la zona del medio.")
st.markdown(verdict(co.price, res))
if model == "FCFF":
    _imp = implied_wacc(co, p, model)
    if _imp is None:
        st.caption("🔁 **DCF inverso:** no se puede calcular (el valor da 0 o el precio queda fuera del rango de búsqueda).")
    else:
        _dif = _imp["wacc_model"] - _imp["wacc"]
        _lec = ("el mercado descuenta menos que este modelo, o espera flujos mejores que los nuestros" if _dif > 0.0025 else
                "el mercado exige más retorno que este modelo, o espera flujos peores que los nuestros" if _dif < -0.0025 else
                "el modelo y el mercado coinciden")
        st.caption(f"🔁 **DCF inverso:** con los flujos de este modelo, el precio actual equivale a un WACC inicial de **{_imp['wacc']:.2%}** "
                   f"(el modelo usa {_imp['wacc_model']:.2%}): {_lec}. No distingue si la diferencia está en la tasa o en los flujos.")
if model in ("FCFF", "REIT") and co.rev_g_y0 is None and co.rev_g_y1 is None and growth_override is None:
    st.warning("**Confianza menor:** sin acceso a estimaciones de analistas, proyecciones basadas en la mediana histórica de crecimiento. "
               f"Podés mirarlas vos en [Yahoo Finance](https://finance.yahoo.com/quote/{co.ticker}/analysis) o en "
               "[Finviz](https://finviz.com/quote.ashx?t=" + co.ticker.split('.')[0] + ") y, si querés, fijar el crecimiento de los "
               "años 1 y 2 en Ajustes avanzados (opcional).")
sens = sensitivity(co, p, model)
st.download_button("⬇️ Descargar resultado en Excel",
                   build_excel(co, res, base, sens, p, model),
                   file_name=f"DCF_{co.ticker}_{datetime.now():%Y%m%d}.xlsx",
                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ----------------------------------------------------------------------------
# Mapa de métodos (v2.48): otras formas de mirar el valor, al lado del DCF. No cambian el valor del modelo.
# ----------------------------------------------------------------------------
if model != "FCFF":
    st.caption("🗺️ **Mapa de métodos:** solo se calcula para empresas valuadas por flujo libre (FCFF).")
elif base.ok:
    with st.expander("🗺️ Mapa de métodos: 5 formas de mirar el valor (no cambian el valor del modelo)", expanded=True):
        _mc1, _mc2 = st.columns(2)
        _dg = _mc1.select_slider("Método 3 (Dorsey): crecimiento anual del flujo de caja libre", options=[0.03, 0.04, 0.05, 0.06, 0.07, 0.08, 0.09, 0.10],
                                 value=DORSEY_G_DEFAULT, format_func=lambda v: f"{v:.0%}", key=K("dorsey_g"),
                                 help="Dorsey, en su ejemplo de Clorox, usa un crecimiento prudente y constante (5 %) por 10 años y una perpetuidad de 3 %.")
        _dr = _mc2.number_input(f"Método 3 (Dorsey): tasa de descuento en % (0 = automática: la mayor entre {DORSEY_R_MIN:.0%} y el costo del capital propio)",
                                0.0, 25.0, 0.0, 0.5, key=K("dorsey_r"))
        mp = build_methods_map(co, base, model, dorsey_g=_dg, dorsey_R=(_dr / 100.0 if _dr > 0 else None))
        if not mp["applicable"]:
            st.info(mp["reason"])
        else:
            _rows = []
            for it in mp["items"]:
                if it["valor"] is None:
                    v_txt, r_txt = "n/d", "n/d"
                elif it["rango"]:
                    v_txt = f"{money(it['valor'], co.currency)}  (rango {money(it['rango'][0], '', 2)} – {money(it['rango'][1], '', 2)})"
                    r_txt = f"{it['ratio']:.2f}×  ({it['ratio_rango'][0]:.2f}–{it['ratio_rango'][1]:.2f}×)"
                else:
                    v_txt, r_txt = money(it["valor"], co.currency), f"{it['ratio']:.2f}×"
                _rows.append({"Método": it["nombre"], "Valor por acción": v_txt, "Valor ÷ precio": r_txt, "Qué es": it["nota"]})
            st.dataframe(pd.DataFrame(_rows).set_index("Método"), **STRETCH)
            _cd = pd.DataFrame([{"Método": it["nombre"].split(" (")[0], "Valor": it["valor"], "Tipo": "circular" if it["circular"] else "independiente"}
                                for it in mp["items"] if it["valor"] is not None])
            if not _cd.empty:
                _bars = alt.Chart(_cd).mark_bar().encode(
                    y=alt.Y("Método:N", sort=None, title=None), x=alt.X("Valor:Q", title=f"Valor por acción ({co.currency})"),
                    color=alt.Color("Tipo:N", legend=None, scale=alt.Scale(domain=["independiente", "circular"], range=["#4f81bd", "#aaaaaa"])))
                _rl = alt.Chart(pd.DataFrame({"x": [co.price]})).mark_rule(strokeDash=[6, 4], color="#444", size=2).encode(x="x:Q")
                st.altair_chart((_bars + _rl).properties(height=40 * len(_cd) + 40), **STRETCH)
                st.caption(f"Línea punteada: precio actual {money(co.price, co.currency)}. Gris: método circular (el valor sale cerca del precio por construcción).")
            if mp["rango"]:
                st.markdown(f"**Rango de los métodos independientes (1 a 4):** {money(mp['rango'][0], co.currency)} – {money(mp['rango'][1], co.currency)} "
                            f"contra un precio de {money(co.price, co.currency)}. No se promedian: cada uno responde una pregunta distinta y tres comparten la tasa de descuento.")
            pn = mp["panel"]
            if pn:
                st.markdown("**Panel a 5 años: ¿qué valdría y qué rendimiento tendría comprar hoy?**")
                _pr = [
                    {"": "Valor por acción en 5 años según el modelo", "Resultado": money(pn["valor5_modelo"], co.currency),
                     "Si comprás hoy → retorno anual": pct(pn["retorno_modelo"]) if pn["retorno_modelo"] is not None else "n/d"},
                ]
                if pn["valor5_mult_hoy"] is not None:
                    _pr.append({"": f"Valor por acción en 5 años si el mercado mantiene el múltiplo de hoy (EV/EBIT {pn['mult_hoy']:.1f}×)",
                                "Resultado": money(pn["valor5_mult_hoy"], co.currency),
                                "Si comprás hoy → retorno anual": pct(pn["retorno_mult_hoy"]) if pn["retorno_mult_hoy"] is not None else "n/d"})
                st.dataframe(pd.DataFrame(_pr).set_index(""), **STRETCH)
                if pn["mult_requerido"] is not None and pn["mult_hoy"]:
                    _mr, _mh = pn["mult_requerido"], pn["mult_hoy"]
                    _lec5 = ("el precio de hoy asume que dentro de 5 años el mercado pagará MÁS por esta empresa que hoy" if _mr > _mh * 1.05 else
                             "alcanza con que el mercado pague menos que hoy: el precio actual deja margen" if _mr < _mh * 0.95 else
                             "el precio de hoy exige que el mercado siga pagando lo mismo que hoy")
                    st.markdown(f"**Múltiplo que hace falta en 5 años** (EV/EBIT) para ganar el costo del capital del modelo ({pn['wacc']:.1%}): "
                                f"**{_mr:.1f}×** · hoy el mercado paga {_mh:.1f}× · el que sale de nuestro modelo es {pn['mult5_model']:.1f}×. → {_lec5}.")
            for _n in mp["notas"]:
                st.caption("• " + _n)
            st.caption("Simplificaciones: se supone que todo el flujo libre llega al accionista y que la deuda neta y las acciones no cambian. "
                       "Los múltiplos de industria son de Damodaran (EE.UU., " + MULT_DATE + "). El método 5 es una aproximación nuestra, NO es el número de Alpha Spread.")

if all_warns:
    with st.expander(f"⚠️ Avisos y controles ({len(all_warns)})", expanded=True):
        for w in all_warns:
            st.write("• " + w)

def _bridge_rows():
    def line(label, fn, fmt=None):
        out = {"": label}
        for n in SCEN_ORDER:
            out[n] = fn(res[n]) if res[n].ok else "n/d"
        return out
    m = lambda v: money(v, "", 0)
    if model == "FCFF":
        return [line("Valor presente de los flujos de los años 1 a N", lambda r: m(r.extras["pv_explicit"])),
                line("+ Valor presente del valor terminal (de ahí en adelante)", lambda r: m(r.extras["pv_tv"])),
                line("= Valor de la empresa (EV)", lambda r: m(r.ev)),
                line("− Deuda", lambda r: m(-r.extras["debt"])),
                line("+ Caja", lambda r: m(r.extras["cash"])),
                line("+ Inversiones en otras empresas (valor contable)", lambda r: m(r.extras["nonop"])),
                line("− Participaciones minoritarias", lambda r: m(-r.extras["minority"])),
                line("= Valor del patrimonio", lambda r: m(r.equity_value)),
                line("÷ Acciones (millones)", lambda r: f"{co.shares_m:,.0f}"),
                line("= Valor por acción", lambda r: money(r.per_share, co.currency)),
                line("Peso del valor terminal en el total", lambda r: pct(r.tv_share, 0))]
    if model == "ER":
        return [line("Patrimonio contable (book value)", lambda r: m(r.extras["bv0"])),
                line("+ Valor presente de los retornos por encima del costo del capital (años 1 a N)", lambda r: m(r.extras["pv_explicit"])),
                line("+ Valor presente del valor terminal", lambda r: m(r.extras["pv_tv"])),
                line("= Valor del patrimonio", lambda r: m(r.equity_value)),
                line("÷ Acciones (millones)", lambda r: f"{co.shares_m:,.0f}"),
                line("= Valor por acción", lambda r: money(r.per_share, co.currency))]
    return [line("Valor presente de las distribuciones (años 1 a N)", lambda r: m(r.extras["pv_explicit"])),
            line("+ Valor presente del valor terminal", lambda r: m(r.extras["pv_tv"])),
            line("= Valor del patrimonio", lambda r: m(r.equity_value)),
            line("÷ Acciones (millones)", lambda r: f"{co.shares_m:,.0f}"),
            line("= Valor por acción", lambda r: money(r.per_share, co.currency))]


with st.expander("🧮 ¿Cómo se llega al valor por acción? (paso a paso)"):
    st.dataframe(pd.DataFrame(_bridge_rows()).set_index(""), **STRETCH)
    st.caption(f"Montos en millones de {co.currency}. N = horizonte elegido. Cada columna es un escenario.")

# ----------------------------------------------------------------------------
# Pestañas
# ----------------------------------------------------------------------------
t1, t2, t3, t4, t5 = st.tabs(["Supuestos por escenario", "Proyección (Base)", "Sensibilidad", "Controles de cordura",
                              "Metodología y límites"])

with t1:
    def row(label, fn):
        return {"": label, **{n: (fn(res[n]) if res[n].ok else "n/d") for n in SCEN_ORDER}}

    rows = []
    if model == "FCFF":
        rows = [
            row("WACC (años 1–5 → terminal)", lambda r: f"{pct(r.extras['wacc'], 2)} → {pct(r.extras['wacc_T'], 2)}"),
            row("Costo del capital propio (Ke)", lambda r: pct(r.extras["ke"], 2)),
            row("Crecimiento ingresos año 1", lambda r: pct(r.extras["g1"])),
            row("Crecimiento ingresos año 2", lambda r: pct(r.extras["g2"])),
            row("g terminal", lambda r: pct(r.g_terminal, 2)),
            row("Margen EBIT actual (TTM)", lambda r: pct(r.extras["margin0"])),
            row("Margen EBIT objetivo", lambda r: pct(r.extras["margin_target"])),
            row("ROIC inicial", lambda r: pct(r.extras["roic0"])),
            row("Foso (moat)", lambda r: f"{r.extras['moat_label']} ({r.extras['moat_src']})"),
            row("ROIC terminal (WACC + spread)", lambda r: pct(r.extras["roicT"])),
            row("Valor empresa (EV, M)", lambda r: money(r.ev, "", 0)),
            row("Caja (M)", lambda r: money(r.extras["cash"], "", 0)),
            row("Deuda (M)", lambda r: money(r.extras["debt"], "", 0)),
            row("Valor por acción", lambda r: money(r.per_share, co.currency)),
        ]
    elif model == "ER":
        rows = [
            row("Costo del capital propio (Ke)", lambda r: pct(r.ke, 2)),
            row("ROE TTM", lambda r: pct(r.extras["roe_ttm"])),
            row("ROE normalizado de partida", lambda r: pct(r.extras["roe_norm"])),
            row("ROE de largo plazo", lambda r: pct(r.extras["roe_T"])),
            row("Persistencia", lambda r: f"{r.extras['persistence']:.2f}"),
            row("Payout total inicial", lambda r: pct(r.extras["payout0"], 0)),
            row("g terminal", lambda r: pct(r.g_terminal, 2)),
            row("Book value por acción", lambda r: money(r.extras["bv_ps"], co.currency)),
            row("P/B implícito", lambda r: f"{r.extras['pb_implied']:.2f}x"),
            row("Valor por acción", lambda r: money(r.per_share, co.currency)),
        ]
    else:
        rows = [
            row("Costo del capital propio (Ke)", lambda r: pct(r.ke, 2)),
            row("AFFO por acción (partida)", lambda r: money(r.extras["affo_ps"], co.currency)),
            row("Crecimiento AFFO inicial", lambda r: pct(r.extras["g0"])),
            row("g terminal", lambda r: pct(r.g_terminal, 2)),
            row("Retorno de lo reinvertido (ROI)", lambda r: pct(r.extras["roi"], 2)),
            row("P/AFFO implícito", lambda r: f"{r.extras['p_affo_implied']:.1f}x"),
            row("Valor por acción", lambda r: money(r.per_share, co.currency)),
        ]
    st.dataframe(color_rows(pd.DataFrame(rows).set_index("")), **STRETCH)
    st.caption(COLOR_LEGEND)
    st.markdown("**Referencias para cada supuesto**")
    def _hist_cagr(vals, years):
        v = [x for x in vals if x and x > 0]
        if len(v) < years + 1:
            return None
        return (v[-1] / v[-1 - years]) ** (1 / years) - 1
    ref_rows = []
    if model == "FCFF":
        mo = margin_options(co)
        h3, h4 = _hist_cagr(co.revenue_hist, 3), _hist_cagr(co.revenue_hist, min(4, max(1, len(co.revenue_hist) - 1)))
        ref_rows += [
            ("Crecimiento de ingresos · histórico (3 años, anual)", pct(h3, 1) if h3 is not None else "n/d"),
            ("Crecimiento de ingresos · analistas (año en curso / siguiente)",
             " / ".join(pct(x, 1) if x is not None else "n/d" for x in (co.rev_g_y0, co.rev_g_y1))),
            ("Crecimiento de ingresos · rango de analistas (mín – máx, año siguiente)",
             f"{pct(co.rev_gl_y1, 1)} – {pct(co.rev_gh_y1, 1)}" if co.rev_gl_y1 is not None and co.rev_gh_y1 is not None else "n/d"),
            ("Crecimiento de ingresos · usado (años 1 y 2, Pesimista / Base / Optimista)",
             " | ".join(f"{pct(res[n].extras['g1'], 1)} y {pct(res[n].extras['g2'], 1)}" for n in SCEN_ORDER if res[n].ok)),
            ("Margen EBIT · actual (TTM)", pct(mo.get("ttm"), 1)),
            ("Margen EBIT · mediana histórica + TTM", pct(mo.get("median"), 1)),
            ("Margen EBIT · según analistas", pct(mo["analysts"], 1) if mo.get("analysts") is not None else "n/d (faltan estimaciones)"),
        ]
        _rf, _ = get_ref(ind_name, co.country)
        if _rf and _rf.get("margin") is not None:
            ref_rows.append((f"Margen EBIT · promedio del sector ({ind_name})", pct(_rf["margin"], 1)))
        ref_rows.append(("Margen EBIT · objetivo usado (Base)", pct(base.extras["margin_target"], 1)))
        if _rf and _rf.get("roc") is not None:
            ref_rows.append((f"Retorno sobre capital (ROC) · promedio del sector ({ind_name})", pct(_rf["roc"], 1)))
        ref_rows.append(("ROIC · actual → terminal usado (Pesimista / Base / Optimista)",
                         f"{pct(base.extras['roic0'], 1)} → " + " | ".join(pct(res[n].extras['roicT'], 1) for n in SCEN_ORDER if res[n].ok)
                         + (f"  ({base.extras['roic_note']})" if base.extras.get('roic_note') else "")))
        _ev_mkt = co.market_cap_m + co.debt - co.cash - (co.non_op_assets if p.include_nonop else 0) + co.minority
        _np0 = base.extras.get("nopat0") or 0
        ref_rows.append(("Múltiplo EV/NOPAT · implícito en el terminal (Pesimista / Base / Optimista)",
                         " | ".join(f"{res[n].extras['tv_nopat']:.1f}x" for n in SCEN_ORDER if res[n].ok)))
        if _np0 > 0:
            ref_rows.append(("Múltiplo EV/NOPAT · actual de mercado (TTM)", f"{_ev_mkt / _np0:.1f}x"))
    elif model == "ER":
        _rf, _ = get_ref(ind_name, co.country)
        ref_rows += [("ROE · TTM", pct(base.extras["roe_ttm"], 1)), ("ROE · mediana histórica", pct(base.extras["roe_med"], 1)),
                     ("ROE · normalizado usado (Base)", pct(base.extras["roe_norm"], 1))]
        if _rf and _rf.get("roe") is not None:
            ref_rows.append((f"ROE · promedio del sector ({ind_name})", pct(_rf["roe"], 1)))
    else:
        ref_rows += [("Crecimiento inicial del AFFO (Base)", pct(base.extras["g0"], 1)),
                     ("Crecimiento de ingresos · analistas (año en curso / siguiente)",
                      " / ".join(pct(x, 1) if x is not None else "n/d" for x in (co.rev_g_y0, co.rev_g_y1)))]
    st.dataframe(pd.DataFrame(ref_rows, columns=["Referencia", "Valor"]).set_index("Referencia"), **STRETCH)
    if model == "FCFF":
        with st.expander(f"🏰 Foso estimado: {base.extras['moat_label']} ({base.extras['moat_src']}) — ver razones"):
            for _w in base.extras["moat_reasons"]:
                st.write("• " + _w)
            st.caption("Estimación propia de la app (no es una calificación de Morningstar): Amplio = los tres criterios, "
                       "con el exceso de retorno obligatorio · Estrecho = solo exceso de retorno, o con uno de los otros (también si el ROIC queda a menos de 1 pp del umbral y cumple los otros dos) · Ninguno = sin exceso de retorno. "
                       "Los umbrales son una primera versión y se van a calibrar con casos reales. Podés forzar el foso en la barra lateral.")
    st.caption("Los escenarios cambian solo palancas del negocio: crecimiento de ingresos, margen y ventaja competitiva "
               "(Pesimista: margen objetivo −15 %, moat −2 pp y sin persistencia en financieras · Optimista: margen +15 %, moat +2 pp). Las tasas de descuento y el "
               "crecimiento terminal no cambian entre escenarios: mirá su efecto en la pestaña Sensibilidad.")

with t2:
    df = pd.DataFrame(base.table)
    fmt = {c: "{:,.1f}" for c in df.columns if c not in ("Año",)}
    for c in ("Crecimiento", "Margen", "Impuesto", "Reinversión", "ROIC", "ROE", "Payout", "WACC"):
        if c in df.columns:
            fmt[c] = "{:.2%}" if c == "WACC" else "{:.1%}"
    fmt["Año"] = "{:d}"
    st.dataframe(df.style.format(fmt), **STRETCH, hide_index=True)
    ex = base.extras
    if model == "FCFF":
        st.write(f"**Valor terminal:** {ex['tv']:,.0f} M (VP {ex['pv_tv']:,.0f} M) · "
                 f"**VP flujos explícitos:** {ex['pv_explicit']:,.0f} M · "
                 f"**EV:** {base.ev:,.0f} M → **Patrimonio:** {base.equity_value:,.0f} M")
    else:
        st.write(f"**Valor del patrimonio:** {base.equity_value:,.0f} M")
    st.caption(f"Montos en millones de {co.currency}. Estados contables originales en {co.fin_currency}.")

with t3:
    s = sens
    if s.get("ok"):
        idx = [f"{'WACC' if model == 'FCFF' else 'Ke'} {x:.2%}" for x in s["rates"]]
        colz = [f"g {x:.2%}" for x in s["gs"]]
        sdf = pd.DataFrame(s["grid"], index=idx, columns=colz)

        def colorize(v):
            if v != v:
                return ""
            return "background-color:#e3f1e3" if v >= co.price else "background-color:#f7e1e0"
        st.write(f"Valor por acción (escenario Base). Verde = por encima del precio actual ({co.price:,.2f}).")
        st.caption("Aquí sí se mueven las tasas: el WACC de la columna izquierda desplaza juntos el de la empresa y el del sector "
                   "(hacia el que converge), y las columnas cambian el crecimiento terminal. Los escenarios no tocan estas tasas.")
        st.dataframe(sdf.style.format("{:,.2f}").map(colorize), **STRETCH)
    else:
        st.info("Sin sensibilidad disponible.")

with t4:
    st.markdown("**Controles por dato importado**")
    _fc = field_checks(co, model, p)
    _fc_df = pd.DataFrame([{"": STATE_ICON[c["Estado"]], "Dato": c["Campo"], "Valor usado": c["Valor"], "Qué controlar": c["Control"]}
                           for c in _fc]).set_index("")
    st.dataframe(_fc_df, **STRETCH)
    _n_av = sum(1 for c in _fc if c["Estado"] == "aviso")
    st.caption(f"✅ razonable · ⚠️ revisar ({_n_av}) · ℹ️ dato faltante con supuesto conocido. "
               "Antes de confiar en un valor, mirá especialmente los ⚠️.")
    ex = base.extras
    lines = []
    if model == "FCFF":
        lines += [
            ("Peso del valor terminal en el EV", pct(base.tv_share, 0), "Alerta si > 75 %"),
            ("EV/EBIT de salida implícito", f"{ex['exit_ev_ebit']:.1f}x", "Razonable: ~8–20x según sector"),
            *([("Deuda de financiera propia excluida", f"{ex['captive_excl']:,.0f} M", "No se resta: sus intereses ya están en el EBIT (v2.41)")] if ex.get('captive_excl') else []),
            ("ROIC terminal vs WACC terminal", f"{pct(ex['roicT'], 1)} vs {pct(ex['wacc_T'], 1)}", f"Foso: {ex['moat_label']} ({ex['moat_src']})" + (f" · premio parcial interno {ex['moat_partial']:.0%}" if ex.get('moat_partial') else "")),
            ("Beta usada / Rf", f"{ex['beta']:.2f} / {pct(ex['rf'], 2)}", "Con datos de industria: promedio entre la beta del sector (re-apalancada) y la propia, piso 0,5 y tope 2,0. Sin industria: Blume 0,67·β + 0,33"),
            ("Beta terminal", f"{ex['beta_T']:.2f}" if ex.get('beta_T') is not None else "n/d", "Promedio entre la beta de los primeros años (con la deuda de la industria) y 1,0: 25 % propia, 25 % industria, 50 % mercado"),
            ("WACC: empresa → sector (terminal)", f"{pct(ex['wacc'], 2)} → {pct(ex['wacc_T'], 2)}", "El costo del capital converge al terminal (deuda de la industria, beta terminal) entre los años 5 y N"),
            ("Costo de la deuda (antes de impuestos)", pct(ex['kd'], 2), ex.get('kd_note', "")),
        ]
    elif model == "ER":
        lines += [
            ("P/B actual del mercado", f"{ex['pb_current']:.2f}x", ""),
            ("P/B implícito (Base)", f"{ex['pb_implied']:.2f}x", "Debe ser coherente con ROE vs Ke"),
            ("P/E implícito (Base)", f"{ex['pe_implied']:.1f}x", ""),
        ]
    else:
        lines += [
            ("P/AFFO actual", f"{ex['p_affo_current']:.1f}x", ""),
            ("P/AFFO implícito (Base)", f"{ex['p_affo_implied']:.1f}x", ""),
            ("Fuente del AFFO", ex["affo_src"], ""),
        ]
    if co.target_price:
        lines.append(("Precio objetivo de analistas (solo referencia)", money(co.target_price, co.currency),
                      f"{co.n_analysts or '?'} analistas"))
    lines.append(("Fuente de datos TTM", co.ttm_source or "n/d", f"al {co.asof}" if co.asof else ""))
    st.dataframe(pd.DataFrame(lines, columns=["Control", "Valor", "Referencia"]).set_index("Control"),
                 **STRETCH)
    # ---- contraste con el sector (Damodaran) ----
    st.markdown("**Contraste con el sector**")
    ref, ref_region = get_ref(ind_name, co.country)
    if not ref:
        st.info("No pude detectar la industria de esta empresa (Yahoo no la informó). "
                "Elegila en la barra lateral → Ajustes avanzados → «Industria para comparar».")
    else:
        def read(emp, sec, tol=0.25, absolute=None):
            if emp is None or sec is None or emp != emp:
                return "n/d"
            diff = (emp - sec) if absolute is not None else (emp / sec - 1 if sec else 0)
            lim = absolute if absolute is not None else tol
            return "similar" if abs(diff) <= lim else ("por encima del sector" if diff > 0 else "por debajo del sector")
        rows = [("Beta usada", f"{ex['beta']:.2f}", f"{ref['beta']:.2f}",
                 read(ex['beta'], ref['beta'], 0.3))]
        if model == "FCFF":
            rows += [("WACC", pct(ex['wacc'], 1), pct(ref['wacc'], 1), read(ex['wacc'], ref['wacc'], absolute=0.01)),
                     ("Margen EBIT actual (TTM)", pct(ex['margin0'], 1), pct(ref['margin'], 1),
                      read(ex['margin0'], ref['margin'], 0.3)),
                     ("Margen EBIT objetivo", pct(ex['margin_target'], 1), pct(ref['margin'], 1),
                      read(ex['margin_target'], ref['margin'], 0.3)),
                     ("Retorno sobre capital hoy (ROIC)", pct(ex['roic0'], 1), pct(ref.get('roc'), 1) if ref.get('roc') is not None else "n/d",
                      read(ex['roic0'], ref.get('roc'), 0.4))]
        elif model == "ER":
            rows += [("Costo del capital propio (Ke)", pct(ex['ke'], 1), pct(ref['ke'], 1), read(ex['ke'], ref['ke'], absolute=0.01)),
                     ("ROE actual (TTM)", pct(ex['roe_ttm'], 1), pct(ref['roe'], 1), read(ex['roe_ttm'], ref['roe'], 0.3)),
                     ("ROE normalizado usado", pct(ex['roe_norm'], 1), pct(ref['roe'], 1), read(ex['roe_norm'], ref['roe'], 0.3))]
        else:
            rows += [("Costo del capital propio (Ke)", pct(ex['ke'], 1), pct(ref['ke'], 1), read(ex['ke'], ref['ke'], absolute=0.01))]
        st.dataframe(pd.DataFrame(rows, columns=["Indicador", "Empresa", f"Sector: {ind_name}", "Lectura"]).set_index("Indicador"),
                     **STRETCH)
        st.caption(f"Promedios de {ind_name} ({ref_region}, {ref['n']} empresas), Damodaran, {IND_DATE}, en dólares. "
                   f"Industria {ind_how}. El WACC/Ke del sector usa la tasa libre de riesgo de enero 2026 (cerca de 4,2 %), "
                   "por eso el de la empresa puede salir algo más alto hoy. Es una referencia para detectar números raros, no un objetivo.")
        if model == "FCFF" and ref.get('margin') and ex['margin_target'] > max(1.5 * ref['margin'], ref['margin'] + 0.05):
            st.warning("El margen objetivo es bastante mayor al típico del sector: revisá si la empresa realmente "
                       "puede sostenerlo (moat) o usá «Mantener TTM» / un margen manual más bajo.")

    for n in SCEN_ORDER:
        if res[n].ok and res[n].warnings:
            with st.expander(f"Avisos del escenario {n}"):
                for w in res[n].warnings:
                    st.write("• " + w)

with t5:
    st.markdown("""
**Qué hace cada modelo**

* **FCFF desde ingresos** (no financieras). Los ingresos de los años 1–2 siguen las estimaciones de analistas
  (si no hay, la mediana histórica acotada, y la app lo marca como «confianza menor»); ese ritmo se **sostiene hasta el año 5** y después
  baja gradualmente hasta *g* terminal en el último año. Pesimista y optimista usan el rango de analistas cuando existe (si no, ±3 pp).
  El margen EBIT converge desde el TTM al objetivo (mediana histórica por defecto; opciones: TTM, analistas, manual). **Tendencia clara (v2.46):** si el margen subió —o bajó— los 4 años seguidos, el objetivo es el punto medio entre la mediana y el margen de hoy (la mediana quedaría dos años atrás); no se aplica en industrias de precio-aceptante (petróleo, minería, acero, químicos básicos, aerolíneas), donde el último margen puede ser un pico o un piso del ciclo, ni si algún año rompe la serie; en los escenarios el objetivo se mueve ∓15 % (relativo). Si el TTM está 8 pp o más bajo el objetivo y los ingresos no caen (posible cargo extraordinario), la proyección arranca del margen normal y no del TTM (y la tasa de impuestos arranca de la marginal del país); el pesimista tampoco se ancla al TTM. Si la empresa
  tiene pérdida operativa, converge a su propio nivel de años rentables (con tope de 1,5× el margen de la industria) o, si nunca fue rentable, al margen promedio de su industria. El impuesto converge de la tasa efectiva a la **marginal del país**
  (tabla de Damodaran) en 5 años. La **reinversión sale solo del crecimiento**: reinversión = *g* / ROIC, con un ROIC inicial normalizado (calculado con el margen objetivo, no con el EBIT del TTM, que puede estar distorsionado por cargos extraordinarios) que converge al
  **ROIC terminal**: con datos de industria, *Amplio* = ROC del sector (Damodaran), *Estrecho* = punto medio con el costo del capital,
  *Ninguno* = costo del capital (nunca por encima del ROIC actual; *Amplio* y *Estrecho* tienen un piso de WACC + 3 / + 2 pp); sin datos, WACC + 0/2/5 pp. Así el crecimiento solo vale si hay
  ventaja competitiva. El **foso se estima automáticamente** con tres criterios (exceso de retorno: ROIC ≥ WACC del sector + 5 pp, obligatorio porque para Damodaran el foso es un retorno sostenidamente por encima del costo del capital; persistencia de los márgenes recientes; y ser mejor que sus pares en margen o retorno). Las productoras de materias primas y aerolíneas quedan limitadas a Estrecho; se puede forzar a mano. Para que no haya un salto brusco en el umbral, una empresa clasificada *Ninguno* que rinde por encima del WACC del sector pero no llega a ese umbral recibe en el valor terminal una fracción del retorno extra de un foso *Estrecho* (proporcional a qué tan cerca está del umbral); es un cálculo interno, y la clasificación que se muestra sigue siendo Amplio / Estrecho / Ninguno (si el premio parcial llega al 50 % o más, se muestra Estrecho, pero el cálculo usa el porcentaje exacto). Si el capital operativo da cero o negativo porque las inversiones no operativas superan al capital invertido, el ROIC se calcula con el capital sin restarlas (se avisa). Si el valor del patrimonio da negativo se muestra 0 (responsabilidad limitada). El **WACC converge** del de la empresa al del sector: constante hasta el año 5 y luego
  lineal hasta el último año; el sector se reconstruye con las mismas tasas (Rf, prima de mercado y país) y la beta y deuda típicas de su industria.
  El capital invertido es operativo (deuda + patrimonio − caja − inversiones en otras empresas). Valor terminal de Gordon sobre
  NOPAT × (1 − g/ROIC) con el WACC del sector; el gasto en **I+D se capitaliza** (se suma al resultado operativo y se amortiza en 4 años, con el I+D no amortizado como activo); si la calificación sintética es peor que BBB el valor se multiplica por (1 − probabilidad de quiebra a 10 años) (tabla del libro de Damodaran); la distancia WACC − g se mantiene en al menos 2,5 pp.
  EV + caja + inversiones en otras empresas − deuda − minoritarios = patrimonio, dividido por las acciones implícitas.
* **Retorno excedente** (bancos, aseguradoras, mREITs, brokers y, desde la v2.40, utilities reguladas de electricidad, gas y agua: el regulador fija lo que ganan sobre su capital y la mayor parte de lo que invierten se financia con deuda, así que el método por EV y flujo libre castigaba de más a todo el grupo; las productoras independientes, renovables y diversificadas siguen por FCFF). Valor = book value + VP de (ROE − Ke) × patrimonio. El ROE
  parte de un valor normalizado y cae hacia Ke + persistencia × (ROE normalizado − Ke); el payout converge al compatible con *g*.
  Si ROE = Ke, el valor es exactamente el book value.
* **AFFO** (REITs). AFFO ≈ (resultado neto + amortización ± resultados por venta de activos) × ratio, o el AFFO/acción que
  ingreses. Lo retenido se reinvierte a Ke + spread, de ahí sale el crecimiento. Se descuenta a Ke.

**Mapa de métodos (v2.48).** Debajo del veredicto, cinco formas de mirar el valor; NINGUNA cambia el valor del modelo y no se promedian. (1) Nuestro DCF a 10 años. (2) El mismo DCF pero cortado en el año 5, con valor final igual al múltiplo EV/EBITDA de la industria (Damodaran, EE.UU.) por el EBITDA del año 5. (3) *Estilo Dorsey*: flujo de caja libre real de Yahoo (flujo operativo − inversión) que crece un % constante 10 años (5 % por defecto, como en su ejemplo de Clorox) y después una perpetuidad de 3 %; tasa de descuento: la mayor entre 9 % y el costo del capital propio del modelo (la podés cambiar). (4) El múltiplo EV/EBITDA de la industria por el EBITDA de la empresa. (5) *Estilo Alpha*: como el 2 pero con el múltiplo EV/EBIT que el mercado paga hoy por la empresa; es circular (el valor sale cerca del precio por construcción) y es una aproximación nuestra, no el número de Alpha Spread. El **panel a 5 años** muestra qué valdría la empresa en 5 años según el modelo, el retorno anual de comprar hoy, y el múltiplo que hace falta en 5 años para ganar el costo del capital. Límites: solo FCFF (no financieras ni REIT); los múltiplos de industria son de EE.UU. aunque la empresa sea de otro país; el método Dorsey puede salir muy bajo si el flujo de caja libre de los últimos 12 meses es transitoriamente chico; se supone que todo el flujo libre llega al accionista.

**Industria de comparación (v2.43).** La app busca cada empresa en la planilla de Damodaran «Companies in each industry» (48.156 empresas, las mismas 94 industrias de sus tablas) por ticker, por sufijo de bolsa o por nombre, y usa su industria exacta; si la empresa no figura, usa la industria de Yahoo traducida con reglas propias. La clasificación de Damodaran es por código industrial (SIC) y a veces no es la intuitiva (Alphabet y Meta en *Software (Entertainment)*; Caterpillar en *Construction Supplies*). Excepción (v2.45): Visa, Mastercard, PayPal, procesadores de pagos, bolsas y proveedores de datos figuran en la planilla junto con los prestamistas; si se valúan por flujo libre y casi no tienen deuda (deuda ≤ 1,5 veces su capitalización), se comparan con *Business & Consumer Services* (donde Damodaran pone a Equifax, TransUnion y Verisk). Los prestamistas (AmEx, Capital One...) se valúan por retorno excedente y conservan su grupo. Se puede cambiar a mano en el selector.

**Costo de capital.** Ke = Rf (bono 10 años EE.UU., spot) + β × prima de mercado (+ prima país). La prima de mercado es la implícita de EE.UU. de Damodaran (3,70 % a oct-2026; se actualiza a mano cada 4-6 meses y la app avisa si tiene más de 6 meses). β = promedio simple entre la beta de la industria (Damodaran; muestra de EE.UU., o la global si la de EE.UU. tiene menos de 20 empresas) re-apalancada con
la deuda de la empresa (el promedio de muchas empresas tiene mucho menos error que la regresión de una sola) y la beta propia de la acción, con piso 0,5 y tope 2,0
(una beta propia ≤ 0 se considera no confiable y se usa solo la de la industria). Sin datos de industria: beta de la acción con
ajuste de Blume (0,67·β + 0,33) y los mismos límites. El WACC terminal, al que converge el de la empresa, usa una beta que es el promedio entre la de los primeros años y 1,0 (25 % propia, 25 % industria, 50 % mercado; en crecimiento estable la beta tiende a la del mercado), desapalancada con la deuda de la empresa y re-apalancada con la deuda típica de la industria, que también da los pesos de capital. Kd = tasa libre de riesgo + spread de calificación sintética
(según cobertura de intereses; tope BBB en utilities y REITs). Ponderado con capitalización y deuda (con arrendamientos si corresponde).
**Crecimiento terminal:** tasa libre de riesgo − 1 punto, entre 3,0 % y 3,5 %. El techo queda por debajo del crecimiento nominal de la economía de EE.UU. (≈ 3,8–4,1 %: PIB real 1,8 % según el CBO más inflación de 2,0–2,3 %), porque una empresa sola no crece para siempre como toda la economía.

**ADRs y acciones extranjeras.** Los estados se convierten a la moneda de cotización a tipo spot. Las acciones se calculan como
capitalización ÷ precio, lo que respeta el ratio del ADR y las clases múltiples. Precios en centavos (GBp, ZAc) se pasan a unidad.

**Límites que conviene conocer**

* Los datos son de Yahoo Finance: pueden tener huecos o errores; los avisos señalan lo que se detectó, no todo lo posible.
* Tasa libre de riesgo en dólares; para monedas con inflación alta, usá mercado en USD o cargá prima país y Rf. El foso automático y sus umbrales son una primera versión: se calibran con casos reales.
* Empresas con pérdida operativa: se valúan suponiendo que el margen llega al promedio de su industria (o al margen manual que cargues); es un escenario, no una proyección, y queda marcado. Con FCF negativo pero EBIT positivo se valúa y se avisa. Cíclicas (energía, materiales): revisar la normalización del margen.
* Redes de pago (V, MA) van por FCFF; prestamistas de tarjetas (AXP, COF) conviene pasarlas a Retorno excedente a mano.
* Gestoras de activos alternativas (BX, KKR) tampoco encajan bien en retorno excedente: probar FCFF.
* El AFFO se estima si no lo cargás; lo ideal es usar el reportado por la compañía.
* Sin ajuste por opciones/dilución futura más allá de las acciones actuales; el EBIT ya incluye el gasto en acciones (SBC).
* Los modelos de financieras y REITs son una primera versión: calibrar con empresas conocidas antes de confiar.
""")
